from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = Path(__file__).resolve().parent
CHARTS = ROOT / "multi_llm_evidence_extraction" / "reports" / "charts"
OUT_FILE = OUT_DIR / "cap_nhat_huong_nghien_cuu_GVHD.pptx"

# Presentation palette: navy/teal for evidence and a restrained amber accent.
NAVY = RGBColor(18, 43, 70)
TEAL = RGBColor(0, 112, 118)
AMBER = RGBColor(196, 117, 27)
INK = RGBColor(28, 38, 48)
MUTED = RGBColor(91, 103, 115)
LINE = RGBColor(215, 222, 228)
PALE_BLUE = RGBColor(237, 246, 250)
PALE_TEAL = RGBColor(232, 246, 244)
PALE_AMBER = RGBColor(252, 246, 234)
WHITE = RGBColor(255, 255, 255)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
LEFT = Inches(0.72)
RIGHT = Inches(12.61)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
blank = prs.slide_layouts[6]


def set_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line if line else fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def textbox(slide, text, x, y, w, h, size=18, color=INK, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
            margin=0.0, line_spacing=1.0):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.space_after = Pt(0)
    p.line_spacing = line_spacing
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=17, color=INK, level=0, gap=7):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.level = level
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
    return box


def title(slide, title_text, kicker=None, number=None):
    textbox(slide, kicker.upper() if kicker else "CẬP NHẬT LUẬN VĂN", LEFT, Inches(0.36), Inches(8.7), Inches(0.25), 10, TEAL, True)
    textbox(slide, title_text, LEFT, Inches(0.66), Inches(11.2), Inches(0.62), 27, NAVY, True)
    rect(slide, LEFT, Inches(1.37), Inches(1.04), Inches(0.05), TEAL)
    if number:
        textbox(slide, f"{number:02d}", Inches(12.14), Inches(0.39), Inches(0.47), Inches(0.3), 11, MUTED, True, align=PP_ALIGN.RIGHT)


def footer(slide, source):
    rect(slide, LEFT, Inches(7.06), Inches(11.89), Inches(0.01), LINE)
    textbox(slide, source, LEFT, Inches(7.14), Inches(10.7), Inches(0.19), 8.3, MUTED)
    textbox(slide, "KLTN | Cập nhật hướng nghiên cứu", Inches(10.85), Inches(7.14), Inches(1.76), Inches(0.19), 8.3, MUTED, align=PP_ALIGN.RIGHT)


def stat_card(slide, x, y, w, h, value, label, tint, accent=TEAL):
    rect(slide, x, y, w, h, tint, tint, True)
    rect(slide, x, y, Inches(0.06), h, accent, accent, True)
    textbox(slide, value, x + Inches(0.24), y + Inches(0.17), w - Inches(0.38), Inches(0.42), 25, NAVY, True)
    textbox(slide, label, x + Inches(0.24), y + Inches(0.72), w - Inches(0.38), h - Inches(0.82), 11.5, MUTED, False)


def add_image(slide, path, x, y, w, h):
    slide.shapes.add_picture(str(path), x, y, width=w, height=h)


def add_note(slide, text):
    notes = slide.notes_slide
    tf = notes.notes_text_frame
    tf.text = text


# Slide 1
slide = prs.slides.add_slide(blank)
set_bg(slide)
rect(slide, Inches(0), Inches(0), Inches(0.23), SLIDE_H, TEAL)
textbox(slide, "CẬP NHẬT HƯỚNG NGHIÊN CỨU LUẬN VĂN", LEFT, Inches(0.58), Inches(7.9), Inches(0.32), 12, TEAL, True)
textbox(slide, "Từ keyword-based\nđến LLM hỗ trợ quyết định", LEFT, Inches(1.08), Inches(8.3), Inches(1.55), 33, NAVY, True, line_spacing=0.93)
textbox(slide, "Mục tiêu mới: kết hợp tín hiệu kỹ thuật, hiểu ngữ cảnh tin tức và bằng chứng có thể kiểm tra.", LEFT, Inches(2.92), Inches(7.8), Inches(0.7), 17, INK)

# Flow blocks
flow_y = Inches(4.38)
blocks = [
    ("Keyword\nđếm từ", "Đã thử nhiều cách\nnhưng không ổn định", PALE_AMBER, AMBER),
    ("LLM\nđọc ngữ cảnh", "Trích relevance, materiality,\ndirection, evidence span", PALE_TEAL, TEAL),
    ("Decision\nsupport", "Evidence card + dashboard\nhỗ trợ nhà đầu tư", PALE_BLUE, NAVY),
]
for idx, (head, body, tint, accent) in enumerate(blocks):
    x = LEFT + Inches(idx * 3.9)
    rect(slide, x, flow_y, Inches(3.32), Inches(1.56), tint, tint, True)
    rect(slide, x, flow_y, Inches(0.06), Inches(1.56), accent, accent, True)
    textbox(slide, head, x + Inches(0.25), flow_y + Inches(0.2), Inches(1.12), Inches(0.7), 18, NAVY, True, line_spacing=0.85)
    textbox(slide, body, x + Inches(1.42), flow_y + Inches(0.28), Inches(1.62), Inches(0.72), 11.2, INK)
    if idx < 2:
        textbox(slide, "→", x + Inches(3.38), flow_y + Inches(0.53), Inches(0.34), Inches(0.34), 24, TEAL, True, align=PP_ALIGN.CENTER)
textbox(slide, "Bản cập nhật ngắn gửi GVHD", LEFT, Inches(6.67), Inches(3.5), Inches(0.3), 11, MUTED)
footer(slide, "Nguồn: advisor_pitch_one_page.md; ket_qua_luan_van_semantic_news_materiality.md")
add_note(slide, "Mở đầu: hướng cũ không bị bỏ đi; được giữ thành baseline phản biện. Hướng mới tập trung hỗ trợ quyết định có bằng chứng, không hứa hẹn tín hiệu mua/bán.")

# Slide 2
slide = prs.slides.add_slide(blank)
set_bg(slide)
title(slide, "Các hạng mục đã hoàn thành", "Tiến độ nghiên cứu", 2)
stat_card(slide, Inches(0.76), Inches(1.75), Inches(2.7), Inches(1.38), "28.062", "bài tin unique\n6 nguồn", PALE_BLUE, NAVY)
stat_card(slide, Inches(3.64), Inches(1.75), Inches(2.7), Inches(1.38), "685.704", "dự báo ML out-of-sample\n3 purged folds", PALE_TEAL, TEAL)
stat_card(slide, Inches(6.52), Inches(1.75), Inches(2.7), Inches(1.38), "150", "LLM consensus rows\n122 eligible", PALE_AMBER, AMBER)
stat_card(slide, Inches(9.40), Inches(1.75), Inches(2.7), Inches(1.38), "25", "evidence packs\n+ decision cards", PALE_BLUE, NAVY)

items_left = [
    "Crawl và chuẩn hóa tin tức; ghép với giá và mã VN30.",
    "Xây dựng đặc trưng kỹ thuật, kiểm định walk-forward có purge 20 phiên.",
    "Xây dựng schema LLM: relevance, materiality, direction, event type, evidence span.",
]
items_right = [
    "Sinh consensus pseudo-label, kiểm tra agreement và QC thủ công.",
    "Tạo evidence pack, decision card, monitoring timeline và dashboard prototype.",
    "Theo dõi outcome tách biệt dữ liệu sẵn có tại thời điểm ra quyết định.",
]
rect(slide, Inches(0.76), Inches(3.55), Inches(5.62), Inches(2.8), WHITE, LINE, True)
rect(slide, Inches(6.62), Inches(3.55), Inches(5.48), Inches(2.8), WHITE, LINE, True)
textbox(slide, "Dữ liệu và mô hình", Inches(1.05), Inches(3.85), Inches(4.8), Inches(0.3), 14, NAVY, True)
bullets(slide, items_left, Inches(1.05), Inches(4.25), Inches(4.95), Inches(1.85), 13)
textbox(slide, "LLM và sản phẩm hỗ trợ", Inches(6.91), Inches(3.85), Inches(4.65), Inches(0.3), 14, NAVY, True)
bullets(slide, items_right, Inches(6.91), Inches(4.25), Inches(4.78), Inches(1.85), 13)
footer(slide, "Nguồn: ket_qua_luan_van_semantic_news_materiality.md; decision-support generated summary")
add_note(slide, "Slide này cho thấy hạ tầng nghiên cứu đã hoàn thành xuyên suốt từ dữ liệu, ML, LLM đến prototype. Không dùng các số này để suy ra hiệu quả đầu tư.")

# Slide 3
slide = prs.slides.add_slide(blank)
set_bg(slide)
title(slide, "Keyword: kết quả phản biện, không phải tín hiệu chính", "Baseline keyword", 3)
textbox(slide, "Keyword có ích để tạo baseline minh bạch. Nhưng rule-based không nắm được ngữ cảnh và materiality.", LEFT, Inches(1.62), Inches(6.6), Inches(0.46), 15, INK)
add_image(slide, CHARTS / "rule_vs_semantic_confusion.png", Inches(7.22), Inches(1.52), Inches(5.27), Inches(4.48))

# Table
x0, y0 = Inches(0.76), Inches(2.45)
cols = [Inches(2.08), Inches(1.18), Inches(1.28)]
headers = ["Đại diện rule", "Accuracy", "Macro-F1"]
rows = [
    ("Direction", "37,7%", "22,8%"),
    ("Event type", "14,8%", "8,7%"),
    ("Materiality", "22,1%", "15,9%"),
    ("Ticker relevance", "68,9%", "20,9%"),
]
for c, header in enumerate(headers):
    cx = x0 + sum(cols[:c])
    rect(slide, cx, y0, cols[c], Inches(0.45), NAVY, NAVY)
    textbox(slide, header, cx + Inches(0.08), y0 + Inches(0.12), cols[c] - Inches(0.16), Inches(0.2), 10.5, WHITE, True, align=PP_ALIGN.CENTER)
for r, row in enumerate(rows):
    yy = y0 + Inches(0.45 + r * 0.48)
    tint = PALE_BLUE if r % 2 == 0 else WHITE
    for c, value in enumerate(row):
        cx = x0 + sum(cols[:c])
        rect(slide, cx, yy, cols[c], Inches(0.48), tint, LINE)
        textbox(slide, value, cx + Inches(0.08), yy + Inches(0.14), cols[c] - Inches(0.16), Inches(0.18), 10.5, INK, c > 0, align=PP_ALIGN.CENTER)

rect(slide, Inches(0.76), Inches(4.72), Inches(5.83), Inches(1.2), PALE_AMBER, PALE_AMBER, True)
textbox(slide, "Nguyên nhân chính", Inches(1.02), Inches(4.95), Inches(2.0), Inches(0.22), 12, AMBER, True)
textbox(slide, "Thiếu context/materiality • ticker mismatch • tin thị trường chung • boilerplate • mixed direction", Inches(1.02), Inches(5.28), Inches(5.2), Inches(0.37), 12, INK)
textbox(slide, "Kết luận: giữ keyword như baseline phản biện; không dùng làm hướng đóng góp chính.", LEFT, Inches(6.25), Inches(11.0), Inches(0.35), 14, NAVY, True)
footer(slide, "Nguồn: ket_qua_luan_van_semantic_news_materiality.md — rule vs semantic, n=122")
add_note(slide, "Nói rõ negative finding là kết quả có giá trị. Keyword không chứng minh tin tức vô ích; nó chỉ cho thấy biểu diễn bằng đếm từ thiếu ngữ cảnh và không đủ tốt cho mục tiêu nghiên cứu.")

# Slide 4
slide = prs.slides.add_slide(blank)
set_bg(slide)
title(slide, "LLM nâng cấp biểu diễn tin tức thành bằng chứng", "Semantic evidence extraction", 4)
textbox(slide, "Thay vì đếm từ, LLM trích xuất cấu trúc và câu bằng chứng có thể truy vết.", LEFT, Inches(1.58), Inches(6.2), Inches(0.38), 15, INK)

schema = [
    ("Ticker relevance", "Tin có liên quan trực tiếp đến mã?"),
    ("Materiality", "Mức độ trọng yếu của thông tin?"),
    ("Direction", "Tác động support / risk / neutral?"),
    ("Event type", "Sự kiện nào đang diễn ra?"),
    ("Evidence span", "Câu nào chứng minh nhãn?"),
]
for i, (head, body) in enumerate(schema):
    y = Inches(2.12 + i * 0.66)
    rect(slide, Inches(0.76), y, Inches(2.15), Inches(0.48), PALE_TEAL, PALE_TEAL, True)
    textbox(slide, head, Inches(0.95), y + Inches(0.14), Inches(1.8), Inches(0.18), 11, TEAL, True, align=PP_ALIGN.CENTER)
    textbox(slide, body, Inches(3.17), y + Inches(0.12), Inches(3.42), Inches(0.25), 12.5, INK)

add_image(slide, CHARTS / "agreement_heatmap.png", Inches(7.06), Inches(1.63), Inches(4.83), Inches(4.45))
stat_card(slide, Inches(7.0), Inches(6.05), Inches(1.6), Inches(0.67), "150", "consensus rows", PALE_BLUE, NAVY)
stat_card(slide, Inches(8.78), Inches(6.05), Inches(1.6), Inches(0.67), "122", "eligible rows", PALE_TEAL, TEAL)
stat_card(slide, Inches(10.56), Inches(6.05), Inches(1.6), Inches(0.67), "0,84", "mean agreement", PALE_AMBER, AMBER)
textbox(slide, "Lưu ý: consensus là pseudo-label, không phải human ground truth.", LEFT, Inches(6.46), Inches(5.9), Inches(0.26), 10.7, MUTED, False)
footer(slide, "Nguồn: ket_qua_luan_van_semantic_news_materiality.md — sections pseudo-label, agreement, manual QC")
add_note(slide, "Nhấn mạnh LLM phục vụ trích xuất có cấu trúc và bằng chứng, không thay thế hoàn toàn đánh giá chuyên gia. Agreement 0,84 là nhất quán giữa annotators, không phải độ chính xác quần thể.")

# Slide 5
slide = prs.slides.add_slide(blank)
set_bg(slide)
title(slide, "Kết quả thực nghiệm: có tín hiệu ban đầu, claim có kiểm soát", "Kết quả và giới hạn", 5)
add_image(slide, CHARTS / "event_window_adjusted_returns.png", Inches(0.76), Inches(1.58), Inches(6.22), Inches(3.62))

rect(slide, Inches(7.30), Inches(1.6), Inches(4.82), Inches(1.38), PALE_TEAL, PALE_TEAL, True)
textbox(slide, "Event-window exploratory", Inches(7.58), Inches(1.86), Inches(3.8), Inches(0.24), 12.5, TEAL, True)
textbox(slide, "1 / 8", Inches(7.58), Inches(2.19), Inches(1.15), Inches(0.38), 25, NAVY, True)
textbox(slide, "test robust-positive sau BH-FDR; materiality cao/trung bình vs thấp tại T+5: +2,23% adjusted return.", Inches(8.71), Inches(2.18), Inches(3.0), Inches(0.44), 10.6, INK)

rect(slide, Inches(7.30), Inches(3.17), Inches(4.82), Inches(1.38), PALE_BLUE, PALE_BLUE, True)
textbox(slide, "Technical + semantic ML", Inches(7.58), Inches(3.43), Inches(3.8), Inches(0.24), 12.5, NAVY, True)
textbox(slide, "3 folds", Inches(7.58), Inches(3.76), Inches(1.35), Inches(0.38), 22, NAVY, True)
textbox(slide, "685.704 dự báo OOS, purge 20 phiên. AUC/BA gần 0,5; chưa có bằng chứng lọc cổ phiếu hữu ích ổn định.", Inches(8.80), Inches(3.73), Inches(2.94), Inches(0.48), 10.6, INK)

rect(slide, Inches(7.30), Inches(4.74), Inches(4.82), Inches(1.10), PALE_AMBER, PALE_AMBER, True)
textbox(slide, "Kỷ luật diễn giải", Inches(7.58), Inches(4.98), Inches(3.6), Inches(0.2), 12.5, AMBER, True)
textbox(slide, "Không claim nhân quả, alpha dài hạn, hay khuyến nghị đầu tư.", Inches(7.58), Inches(5.30), Inches(3.95), Inches(0.26), 11.6, INK)
textbox(slide, "Giá trị hiện tại: bằng chứng exploratory + quy trình kiểm định có thể audit.", LEFT, Inches(5.62), Inches(6.3), Inches(0.42), 13, NAVY, True)
footer(slide, "Nguồn: event_window_stat_tests_report.md; ket_qua_luan_van_semantic_news_materiality.md")
add_note(slide, "Không dùng kết quả này làm lời hứa lợi nhuận. Điểm mạnh là có test correction, confidence interval, purged OOS và kết luận null/weak được báo cáo minh bạch.")

# Slide 6
slide = prs.slides.add_slide(blank)
set_bg(slide)
title(slide, "Prototype hỗ trợ nhà đầu tư và bước tiếp theo", "Định hướng mới", 6)
add_image(slide, ROOT / "streamlit_dashboard.png", Inches(0.76), Inches(1.58), Inches(5.56), Inches(4.63))
rect(slide, Inches(0.76), Inches(6.28), Inches(5.56), Inches(0.47), PALE_BLUE, PALE_BLUE, True)
textbox(slide, "Dashboard: tín hiệu kỹ thuật + evidence semantic + monitoring + outcome review", Inches(0.98), Inches(6.43), Inches(5.07), Inches(0.18), 10.4, NAVY, True, align=PP_ALIGN.CENTER)

textbox(slide, "Giá trị của hướng mới", Inches(6.72), Inches(1.70), Inches(4.85), Inches(0.28), 15, NAVY, True)
bullets(slide, [
    "Không chỉ trả về score: cho thấy tin nào, vì sao trọng yếu và câu bằng chứng.",
    "Tách point-in-time evidence khỏi outcome review để hạn chế hindsight bias.",
    "Keyword giữ vai trò baseline minh bạch và phần phản biện học thuật.",
], Inches(6.72), Inches(2.14), Inches(5.22), Inches(1.63), 13.2)

textbox(slide, "Bước tiếp theo", Inches(6.72), Inches(4.03), Inches(4.85), Inches(0.28), 15, NAVY, True)
next_steps = [
    ("01", "Mở rộng human annotation và kiểm tra liên-chủ thể."),
    ("02", "Audit contamination, ticker matching và dữ liệu point-in-time."),
    ("03", "Paper-trading trước bất kỳ claim nào về tính hữu ích đầu tư."),
]
for i, (num, label) in enumerate(next_steps):
    yy = Inches(4.47 + i * 0.59)
    rect(slide, Inches(6.72), yy, Inches(0.46), Inches(0.38), TEAL, TEAL, True)
    textbox(slide, num, Inches(6.72), yy + Inches(0.10), Inches(0.46), Inches(0.15), 9, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(slide, label, Inches(7.36), yy + Inches(0.10), Inches(4.78), Inches(0.2), 12.2, INK)

footer(slide, "Nguồn: EvidenceTrace dashboard; reports/decision_support; canonical semantic news materiality report")
add_note(slide, "Kết: đề tài chuyển từ tối ưu keyword sang hệ thống hỗ trợ quyết định có bằng chứng. Ba bước sau là điều kiện để tăng độ tin cậy trước bất kỳ kết luận đầu tư nào.")

prs.save(OUT_FILE)
print(OUT_FILE)
